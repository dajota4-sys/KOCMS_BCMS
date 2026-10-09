"""data/*.json -> outputs/BCMS_연계성_V3.xlsx, outputs/BCMS_연계성_V3_근거정의서.md, survey/survey_items_v3.(json|csv)
사전 조건: scripts/verify_quotes.py 실행(outputs/quote_verification.json), scripts/mkmap.py 실행(지도 PNG).
엑셀 수식 재계산은 별도로 LibreOffice recalc(README 참조)."""
import csv,json,os,sys
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from common import load,out,ROOT,DOMAIN_ORDER,CORE_NAME
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image

Q={q['id']:q for q in load('quotes.json')}; LK=load('links.json'); IT=load('items.json'); DF=load('definitions.json')
VB=load('relation_verbs.json'); HL=load('held_links.json'); OB=load('objects.json'); RESP=load('respondent_items.json')
DM=load('dimensions.json'); SRC=load('sources.json')
OBJ={o['id']:o for o in OB}; VBD={v['id']:v for v in VB}; DMD={d['id']:d for d in DM}
ver=json.load(open(out('quote_verification.json'),encoding='utf-8'))
VOK={r['id']:r['ok'] for r in ver['quotes']}; VDATE=ver['date']
assert not ver['failures'],'인용 검증 실패가 있으므로 빌드를 중단합니다'
NQ=len(Q); NL=len(LK); NI=len(IT); NS=len(SRC)
DORD=[d['id'] for d in DM]  # CO TR EV AC TM FC
def nm(i): return CORE_NAME if i=='CORE' else f"{i} {OBJ[i]['name']}"
FN='맑은 고딕'
F=lambda b=False,c='000000',s=10,i=False: Font(name=FN,bold=b,color=c,size=s,italic=i)
HF=PatternFill('solid',fgColor='1D4E89'); INF=PatternFill('solid',fgColor='FFF9DB'); GREY=PatternFill('solid',fgColor='F2F4F7')
OPF=PatternFill('solid',fgColor='FDECEA'); QF=PatternFill('solid',fgColor='EAF3FB'); BLUE='0000FF'
thin=Side(style='thin',color='D0D5DD'); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
WR=Alignment(wrap_text=True,vertical='top'); CT=Alignment(horizontal='center',vertical='center',wrap_text=True)
wb=Workbook()
def hdr(ws,row,vals,widths=None):
    for j,v in enumerate(vals):
        c=ws.cell(row,1+j,v); c.font=F(True,'FFFFFF'); c.fill=HF; c.alignment=CT; c.border=BD
    if widths:
        for j,w in enumerate(widths): ws.column_dimensions[CL(1+j)].width=w
def put(ws,r,c,v,font=None,al=WR,fill=None,fmt=None):
    x=ws.cell(r,c,v); x.font=font or F(); x.alignment=al; x.border=BD
    if fill: x.fill=fill
    if fmt: x.number_format=fmt
    return x
def title(ws,t,sub=None):
    ws['A1']=t; ws['A1'].font=F(True,'101828',14)
    if sub: ws['A2']=sub; ws['A2'].font=F(False,'475467',10,True)
def qtxt(ids): return ' / '.join(ids) if ids else '(없음)'
QR=f'$A$5:$A${4+NQ}'; QE=f'$E$5:$E${4+NQ}'  # 인용근거 ranges
LR1,LR2=5,4+NL                                # 연결 rows
nvsrc=sum(1 for s in SRC if s['verbatim'])

# ---------- 안내
ws=wb.active; ws.title='안내'
title(ws,'BCMS 프로세스 연결 품질 — 문항은행 V3 (객체층·관계층·품질층)','박사논문 프로젝트 · 근거: Arias-Aranda et al.(2026) / Białas(2010) + 외부 학술 문헌 후보 12건(서지만 확인)')
rows=[
('V3 핵심 변경','V2의 "반영·확인·갱신"을 폐기하고, 첨부 문서(연구자 제시)의 3층 모형을 따랐다. 객체층(프로세스 22개) → 관계층(관계동사 11개) → 품질층(정합성·추적성·증빙성·책임성·최신성·환류폐쇄성 6개 차원). 문항은 "객체쌍 + 관계동사 + 품질조건" 구조다.'),
('표기 규칙','파란 열 = 논문 인용(Arias·Białas는 PDF 원문과 자동 대조 완료) · 빨간 열 = 연구자 의견(해석·보완·가정·문항화). 외부 학술 문헌은 서지만 확인했으며 원문 문장은 확인하지 못했다(학술근거 시트).'),
('품질 차원의 근거 강도','두 논문 안에서의 근거(T1): 환류폐쇄성 강 · 정합성·증빙성·책임성·최신성 중 · 추적성 없음. 외부 문헌(T2)은 모든 차원에 후보가 있으나 원문 확인이 필요하다. 상세는 품질차원 시트.'),
('증거유형','연결의 증거유형: E1 직접 26개 / E2 일부 보완 2개 / E3 가정 3개. 품질 차원의 근거 등급(T1/T2/T3)과는 별개다.'),
('시트 구성','정의서 · 품질차원 · 학술근거 · 인용근거 · 연결 · 적용매트릭스 · 관계동사 · 보류연결 · 객체 · 트리플 · 관계매트릭스 · 온톨로지맵 · 문항 · 보유·응답자정보 · 응답입력 · 연결점수 · 차원점수 · 점수요약 · 검증'),
('사용 순서','① 품질차원·학술근거 시트에서 근거 확인(학술근거 시트 노란 칸에 원문 확인 결과 기입) → ② 연결 시트에서 연구자 의견 검토 → ③ 문항 시트에서 전문가 6인 적합도(1~4) 입력 → ④ 응답입력 7행부터 설문 응답 입력 → ⑤ 연결점수·차원점수·점수요약·검증 확인'),
('입력 칸','노란 바탕 + 파란 글씨 = 직접 입력. 응답입력 6행은 예시이며 집계에서 제외(집계는 7~56행).'),
('응답 척도','1 전혀 그렇지 않다 · 2 그렇지 않다 · 3 보통이다 · 4 그렇다 · 5 매우 그렇다 · 9 해당 없음/모르겠음(결측 처리)'),
('문항 수','연결 문항 163개(연결 31개 × 정합성·추적성·증빙성·책임성·최신성 5개 + 환류 연결 8개 × 환류폐쇄성 1개) + 종합 7 + 검증 9 = 179문항. 응답자 1인이 모두 답하기는 어려우므로 영역별 분할 배포를 권장한다.'),
('중요한 한계','연결 31개는 전문가 검증 전 가설이다. 6개 차원과 "연결 품질" 용어는 연구자가 제시한 초기 가설이며 두 논문에 없다. 외부 문헌의 원문 문장은 이 환경에서 열람할 수 없어 확인하지 못했다.'),
('검증 기록',f'Arias·Białas 인용문 {ver["quotes_ok"]}/{ver["quotes_total"]}, 객체 {ver["objects_ok"]}/{ver["objects_total"]}가 PDF 원문과 일치(검증일 {VDATE}). 외부 문헌 {NS}건은 서지만 확인(원문 대조 {nvsrc}건).'),
]
for i,(a,b) in enumerate(rows):
    put(ws,4+i,1,a,F(True),fill=GREY); put(ws,4+i,2,b); ws.row_dimensions[4+i].height=64 if len(b)>150 else 36
ws.column_dimensions['A'].width=20; ws.column_dimensions['B'].width=140

# ---------- 정의서
w=wb.create_sheet('정의서'); title(w,'정의서 — 용어별 정의, 유형(인용 정의 / 연구자 정의), 근거 인용','유형이 "연구자 정의"인 항목은 두 논문에 없는 연구자의 판단이며, 박사논문에서 연구자가 직접 확정해야 하는 부분')
hdr(w,4,['ID','용어','유형','정의','근거 인용 ID','근거 인용 원문(첫 번째)','근거·한계 설명'],[7,26,22,70,16,60,60])
for i,d in enumerate(DF):
    r=5+i; res=d['type'].startswith('연구자')
    put(w,r,1,d['id']); put(w,r,2,d['term'],F(True)); put(w,r,3,d['type'],fill=OPF if res else None)
    put(w,r,4,d['definition'],fill=OPF if res else None); put(w,r,5,qtxt(d['quotes']))
    put(w,r,6,f'=IFERROR(INDEX(인용근거!{QE},MATCH(LEFT(E{r},3),인용근거!{QR},0)),"")' if d['quotes'] else '(해당 인용 없음)',fill=QF if d['quotes'] else None)
    put(w,r,7,d['note']); w.row_dimensions[r].height=95
w.freeze_panes='C5'

# ---------- 품질차원
w=wb.create_sheet('품질차원'); title(w,'품질층 6개 차원 — 정의와 근거 (T1 두 논문 원문 / T2 외부 문헌 후보 / T3 연구자 정의)','첨부 문서(연구자 제시)의 6차원을 초기 가설로 사용. 근거 강도는 두 논문 안에서의 직접성을 뜻한다.')
hdr(w,4,['순서','코드','차원','영문','첨부 문서의 정의 (연구자 제시)','본 연구의 조작적 정의 (연구자)','문항 틀','적용 연결 수','T1 인용 ID (원문 대조 완료)','T1 대표 인용 원문','T1 근거 강도','T1 근거 설명','T2 외부 문헌 후보 ID','T2 외부 문헌 설명','다른 차원과의 겹침 (연구자 우려)'],[6,7,12,16,36,44,46,9,18,60,9,60,16,50,40])
for i,d in enumerate(DM):
    r=5+i; col=CL(2+i)   # 적용매트릭스 columns B..G
    first=(d['t1'] or d.get('t1_weak') or [''])[0]
    vals=[d['order'],d['id'],d['name'],d['en'],d['attachment_def'],d['researcher_def'],d['item_form']]
    for j,v in enumerate(vals): put(w,r,1+j,v,F(True) if j==2 else None,fill=OPF if j in(5,6) else None)
    put(w,r,8,f'=SUM(적용매트릭스!{col}5:{col}{4+NL})',al=CT)
    put(w,r,9,qtxt(d['t1'])+((' (약한 간접: '+', '.join(d['t1_weak'])+')') if d.get('t1_weak') else ''))
    put(w,r,10,f'=IFERROR(INDEX(인용근거!{QE},MATCH("{first}",인용근거!{QR},0)),"두 논문에 직접 근거 없음")',fill=QF)
    st=d['t1_strength']; put(w,r,11,st,F(True,{'강':'2E9E6B','중':'B54708','없음':'D64545'}[st]),CT)
    put(w,r,12,d['t1_note'],fill=QF); put(w,r,13,qtxt(d['ext'])); put(w,r,14,d['ext_note']); put(w,r,15,d['overlap'],fill=OPF)
    w.row_dimensions[r].height=150
w.freeze_panes='D5'

# ---------- 학술근거
w=wb.create_sheet('학술근거'); title(w,'외부 학술 근거 후보 12건 — 서지는 확인, 원문 문장은 미확인','네트워크 정책(egress 차단)으로 원문을 열람하지 못했다. 연구자가 원문에서 정의·문장을 확인해 노란 칸에 쪽 번호와 원문을 기입하면 근거가 확정된다.')
hdr(w,4,['ID','유형','서지(APA 형식)','연도','품질 차원','요지 (검색 결과 요약 기준, 인용 아님)','서지 확인','원문 확인 상태','DOI / URL','원문 쪽 번호 (연구자 기입)','원문 문장 (연구자 기입)','확인 완료 (Y/N)'],[6,12,70,7,16,60,22,34,36,14,60,12])
for i,s in enumerate(SRC):
    r=5+i
    dm=', '.join(DMD[d]['name'] for d in s['dims']) if s['dims'] else '방법(구성·검증)'
    link=s['doi'] and ('doi:'+s['doi']) or s['url'] or '(서지만)'
    if s['doi'] and s['url']: link='doi:'+s['doi']+' / '+s['url']
    for j,v in enumerate([s['id'],s['type'],s['citation'],s['year'],dm,s['claim'],s['bib_status'],s['text_status'],link]): put(w,r,1+j,v,F(True) if j==0 else None)
    for j in range(3): put(w,r,10+j,None,F(c=BLUE),CT if j!=1 else WR,INF)
    w.row_dimensions[r].height=78
w.freeze_panes='D5'

# ---------- 인용근거
w=wb.create_sheet('인용근거'); title(w,f'인용 근거 라이브러리 — Arias·Białas 원문과 자동 대조한 인용문 {NQ}개','쪽 = PDF 쪽 번호(Arias는 저널 쪽 번호와 동일). 한국어 번역은 연구자 번역(참고용)이며 대조 대상이 아님.')
hdr(w,4,['인용ID','출처','PDF 쪽','절','원문(영어, verbatim)','한국어 번역(연구자 번역·참고)','사용처','원문 대조'],[8,9,8,22,80,60,40,18])
for i,q in enumerate(Q.values()):
    r=5+i
    for j,v in enumerate([q['id'],q['source'],q['pages'],q['section'],q['text'],q['ko'],q['use'],f"일치 ({VDATE})" if VOK[q['id']] else '불일치']): put(w,r,1+j,v,F(True) if j==0 else None,fill=QF if j==4 else None)
    w.row_dimensions[r].height=62
w.freeze_panes='E5'; w.auto_filter.ref=f'A4:H{4+NQ}'

# ---------- 연결
wl=wb.create_sheet('연결'); title(wl,f'연결 {NL}개 — 인용 근거(파랑)와 연구자 의견(빨강)을 분리 · 연결근거 매트릭스','E1 직접 / E2 부분(연구자 보완) / E3 가정. 주 인용 원문은 인용근거 시트에서 자동으로 가져옴.')
H=['ID','영역','연결(S→T)','출발 S','도착 T','관계동사 ID','관계동사','증거유형','집단','환류 연결','주 인용 ID','주 인용 원문(PDF 대조 완료)','보조 인용 ID','인용이 확립하는 사실','연구자 의견: 해석·보완·가정','연구자 의견: 품질 차원 적용','Białas 대응','검증 필요 사항','문항 수','기대 문항 수','점검']
hdr(wl,4,H,[6,18,26,8,8,8,14,9,6,8,10,70,14,55,60,36,35,35,8,8,8])
for i,l in enumerate(LK):
    r=5+i
    vals=[l['id'],l['domain'],f"{nm(l['source'])} → {nm(l['target'])}",l['source'],l['target'],l['verb'],None,l['evidence'],'집단' if l['group'] else '','Y' if 'FC' in l['dims'] else '',qtxt(l['quotes_primary'])]
    for j,v in enumerate(vals): put(wl,r,1+j,v,F(True) if j in(0,7) else None,CT if j in(0,3,4,5,7,8,9) else WR)
    put(wl,r,7,f'=IFERROR(INDEX(관계동사!$B$5:$B${4+len(VB)},MATCH(F{r},관계동사!$A$5:$A${4+len(VB)},0)),"")')
    wl.cell(r,8).fill=PatternFill('solid',fgColor={'E1':'D1FADF','E2':'FEF0C7','E3':'FEE4E2'}[l['evidence']])
    put(wl,r,12,f'=IFERROR(INDEX(인용근거!{QE},MATCH(LEFT(K{r},3),인용근거!{QR},0)),"원문에 S→T 연결 서술 없음")',fill=QF)
    put(wl,r,13,qtxt(l['quotes_support'])); put(wl,r,14,l['established'],fill=QF)
    put(wl,r,15,l['researcher'],fill=OPF); put(wl,r,16,l['quality_note'],fill=OPF); put(wl,r,17,l['bialas']); put(wl,r,18,l['validation'])
    put(wl,r,19,f'=COUNTIF(문항!$D$5:$D${4+NI},A{r})',al=CT); put(wl,r,20,f'=SUM(적용매트릭스!B{r}:G{r})',al=CT); put(wl,r,21,f'=IF(S{r}=T{r},"OK","확인")',F(True),CT)
    wl.row_dimensions[r].height=150
wl.freeze_panes='D5'; wl.auto_filter.ref=f'A4:U{4+NL}'

# ---------- 적용매트릭스
w=wb.create_sheet('적용매트릭스'); title(w,'연결 × 품질 차원 적용 매트릭스 (1 = 해당 차원 문항 있음)','정합성·추적성·증빙성·책임성·최신성은 전 연결에 적용, 환류폐쇄성은 환류 연결 8개에만 적용(연구자 분류). 이 표가 문항 수를 결정한다.')
hdr(w,4,['연결 ID']+[d['name'] for d in DM]+['합계','연결(S→T)'],[10,10,10,10,10,10,12,8,40])
for i,l in enumerate(LK):
    r=5+i; put(w,r,1,l['id'],F(True),CT)
    for j,d in enumerate(DORD): put(w,r,2+j,1 if d in l['dims'] else 0,al=CT,fill=PatternFill('solid',fgColor='D6E6F7') if d in l['dims'] else None)
    put(w,r,8,f'=SUM(B{r}:G{r})',F(True),CT); put(w,r,9,f"{l['source']} → {l['target']}")
r=5+NL; put(w,r,1,'합계',F(True),CT)
for j in range(6): put(w,r,2+j,f'=SUM({CL(2+j)}5:{CL(2+j)}{4+NL})',F(True),CT)
put(w,r,8,f'=SUM(H5:H{4+NL})',F(True),CT)
w.freeze_panes='B5'

# ---------- 관계동사
w=wb.create_sheet('관계동사'); title(w,'관계층 — 관계동사 11개 (연구자 정의, Arias 서술문에서 동사 추출)','첨부 문서의 동사 중 일치·추적·증빙은 품질층 차원과 겹쳐 품질층으로 옮겼다. V10(승인한다)은 문헌에서 확인되나 이번 연결 31개에는 쓰지 않았다.')
hdr(w,4,['ID','관계동사','영문 predicate','의미 (S→T)','판정 질문','Arias 원문 동사 (영문 표현)','근거 인용 ID','첨부 문서 동사와의 대응','사용 연결','사용 수'],[6,18,18,46,40,44,14,36,36,8])
for i,v in enumerate(VB):
    r=5+i
    for j,x in enumerate([v['id'],v['verb'],v['predicate'],v['definition'],v['test'],v['arias_verbs'],qtxt(v['quotes']),v['attachment_map'],', '.join(v['links']) or '(미사용)']): put(w,r,1+j,x,F(True) if j==1 else None,fill=OPF if j in(3,4) else QF if j==5 else None)
    put(w,r,10,f'=COUNTIF(연결!$F${LR1}:$F${LR2},A{r})',al=CT); w.row_dimensions[r].height=48

# ---------- 보류연결
w=wb.create_sheet('보류연결'); title(w,'보류·제외한 후보 연결 — 채택하지 않은 이유','연구자 결정란에 포함/제외를 직접 기록(노란 칸)')
hdr(w,4,['ID','출발','도착','후보 내용','채택하지 않은 이유(인용 근거 부재)','권고','연구자 결정'],[7,10,12,38,70,16,18])
for i,h in enumerate(HL):
    for j,v in enumerate([h['id'],h['source'],h['target'],h['idea'],h['reason'],h['decision']]): put(w,5+i,1+j,v)
    put(w,5+i,7,None,F(c=BLUE),CT,INF); w.row_dimensions[5+i].height=48

# ---------- 객체
wo=wb.create_sheet('객체'); title(wo,'객체층 — Arias Table 5의 22개 BCMS 프로세스','범주·PDCA·재난관리주기: Table 5(p.14) / 전문가 지명률: Table 2–4(pp.11–13, 전문가 39명) / 대표 산출물: 연구자 제안')
hdr(wo,4,['ID','프로세스(한글)','Arias 원문 이름','범주','PDCA','재난관리주기','전문가 지명률(%)','Arias 합의 구분(80/20)','대표 산출물(연구자 제안)','관련 연결 수'],[8,26,50,8,9,14,13,22,36,11])
for i,o in enumerate(OB):
    r=5+i
    for j,v in enumerate([o['id'],o['name'],o['arias_name'],o['category'],o['pdca'],o['lifecycle']]): put(wo,r,1+j,v,fill=QF if j in(2,3,4,5) else None)
    put(wo,r,7,o['expert_pct'],F(c=BLUE),CT,QF)
    put(wo,r,8,f'=IF(G{r}>=80,"범주 1 (채택)",IF(G{r}>=20,"범주 2 (토론)","범주 3 (기각)"))'); put(wo,r,9,o['artefact'],fill=OPF)
    put(wo,r,10,f'=COUNTIF(연결!$D${LR1}:$D${LR2},A{r})+COUNTIF(연결!$E${LR1}:$E${LR2},A{r})',al=CT)
r=5+len(OB)+1
put(wo,r,1,'주: 지명률 합의 구분은 참고용이다. Arias는 합의에 미달한 프로세스도 ISO 22301 요건(A45)을 이유로 핵심 프로세스로 유지했고, CRM(3%)도 참조모델에 포함해야 한다고 보았다(p.13). "BCMS 계획"은 프로젝트로 분류되어 제외했다(A48).',F(i=True,c='475467'))
wo.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10); wo.row_dimensions[r].height=36; wo.freeze_panes='C5'

# ---------- 트리플
wt=wb.create_sheet('트리플'); title(wt,'온톨로지 트리플(주어–관계–목적어)','S—predicate→T 방향으로 통일. ProcessLink 트리플은 연결을 개체로 보고 6개 품질 차원을 속성 점수로 연결(Białas의 data slot 응용, 연구자 설계).')
hdr(wt,4,['주어','관계(predicate)','목적어','관계동사','연결 ID','증거유형','구분'],[24,20,24,16,9,9,20]); r=5
for l in LK:
    v=VBD[l['verb']]
    for j,x in enumerate([l['source'],v['predicate'],l['target'],v['verb'],l['id'],l['evidence'],'프로세스 관계']): put(wt,r,1+j,x)
    r+=1
for l in LK:
    tr=[(f"Link_{l['id']}",'rdf:type','ProcessLink'),(f"Link_{l['id']}",'hasSource',l['source']),(f"Link_{l['id']}",'hasTarget',l['target']),(f"Link_{l['id']}",'hasRelationVerb',VBD[l['verb']]['predicate'])]
    tr+=[(f"Link_{l['id']}",'hasQuality_'+DMD[d]['en'].split()[0],f"{l['id']}-{d}") for d in l['dims']]
    for s_,p_,o_ in tr:
        for j,x in enumerate([s_,p_,o_,'',l['id'],l['evidence'],'ProcessLink 객체']): put(wt,r,1+j,x)
        r+=1
wt.freeze_panes='A5'; wt.auto_filter.ref=f'A4:G{r-1}'

# ---------- 매트릭스
wm=wb.create_sheet('관계매트릭스'); title(wm,'관계 매트릭스 — 행(출발 S) → 열(도착 T)','셀 = 연결 ID, 굵은 글씨 = E1, 이탤릭 = E2, 빨간 글씨 = E3. CORE = 집단 노드(연구자 생성)')
ids=[o['id'] for o in OB]+['CORE']
for j,k in enumerate(ids):
    c=wm.cell(4,2+j,k); c.font=F(True,'FFFFFF'); c.fill=HF; c.alignment=CT; c.border=BD; wm.column_dimensions[CL(2+j)].width=8
    c=wm.cell(5+j,1,k); c.font=F(True,'FFFFFF'); c.fill=HF; c.alignment=CT; c.border=BD
wm.column_dimensions['A'].width=8; cell={}
for l in LK: cell.setdefault((l['source'],l['target']),[]).append(l)
for i,a in enumerate(ids):
    for j,b in enumerate(ids):
        ls=cell.get((a,b),[])
        if ls:
            ev=ls[0]['evidence']; put(wm,5+i,2+j,','.join(x['id'] for x in ls),F(ev=='E1','B42318' if ev=='E3' else '000000',9,ev=='E2'),CT,PatternFill('solid',fgColor='D6E6F7'))
        else: put(wm,5+i,2+j,None,al=CT)
wm.freeze_panes='B5'

# ---------- 지도
wi=wb.create_sheet('온톨로지맵'); title(wi,f'온톨로지 맵 — 22개 프로세스와 {NL}개 연결','그림 파일: outputs/BCMS_온톨로지맵_V3.png · 선 모양 = 증거유형(실선 E1 / 파선 E2 / 점선 E3)')
img=Image(out('BCMS_온톨로지맵_V3.png')); img.width=1500; img.height=int(1500*2100/2800); wi.add_image(img,'A4')

# ---------- 문항
wq=wb.create_sheet('문항'); title(wq,f'문항은행 {NI}문항 (연결 {sum(len(l["dims"]) for l in LK)} + 종합 7 + 검증 9) — 객체쌍 + 관계동사 + 품질조건','연결 문항은 모두 "연구자 문항화"이며 연결의 근거는 연결 시트의 인용을 따른다. 전문가 6인의 적합도(1~4)를 M~R열에 입력.')
hdr(wq,4,['문항ID','파트','영역','연결ID','연결(S→T)','관계동사','품질 차원','문항','증거유형','역문항','응답입력 열','근거 구분','전문가1','전문가2','전문가3','전문가4','전문가5','전문가6','I-CVI','판정','추적성 점검'],[10,14,20,8,30,14,16,70,9,8,10,24,8,8,8,8,8,8,9,12,12])
resp_start=2+10+22; col_of={it['id']:CL(resp_start+i) for i,it in enumerate(IT)}
LKD={l['id']:l for l in LK}
for i,it in enumerate(IT):
    r=5+i; lk=LKD.get(it['link'])
    dm=lk['domain'] if lk else ('종합' if it['part'].startswith('D') else '검증')
    put(wq,r,1,it['id'],F(True)); put(wq,r,2,it['part']); put(wq,r,3,dm); put(wq,r,4,it['link'] or None,al=CT)
    put(wq,r,5,f'=IFERROR(INDEX(연결!$C${LR1}:$C${LR2},MATCH(D{r},연결!$A${LR1}:$A${LR2},0)),"")')
    put(wq,r,6,f'=IFERROR(INDEX(연결!$G${LR1}:$G${LR2},MATCH(D{r},연결!$A${LR1}:$A${LR2},0)),"")',al=CT)
    put(wq,r,7,it['state'],al=CT); put(wq,r,8,it['text'],fill=OPF)
    put(wq,r,9,f'=IFERROR(INDEX(연결!$H${LR1}:$H${LR2},MATCH(D{r},연결!$A${LR1}:$A${LR2},0)),"")',al=CT)
    put(wq,r,10,'Y' if it['reverse'] else '',al=CT); put(wq,r,11,col_of[it['id']],al=CT)
    put(wq,r,12,it['basis_type']+('' if not it['basis_quotes'] else ' · 인용 '+qtxt(it['basis_quotes'])))
    for j in range(6): put(wq,r,13+j,None,F(c=BLUE),CT,INF)
    put(wq,r,19,f'=IFERROR(COUNTIF(M{r}:R{r},">=3")/COUNT(M{r}:R{r}),"")',al=CT,fmt='0.00')
    put(wq,r,20,f'=IF(S{r}="","",IF(S{r}>=검증!$B$4,"채택","수정·토론"))',al=CT)
    put(wq,r,21,f'=IF(COUNTIF($A$5:$A${4+NI},A{r})<>1,"ID 중복",IF(B{r}="C 연결",IF(COUNTIF(연결!$A${LR1}:$A${LR2},D{r})=1,"OK","연결 없음"),"OK"))',al=CT)
    wq.row_dimensions[r].height=34
wq.freeze_panes='I5'; wq.auto_filter.ref=f'A4:U{4+NI}'
dv=DataValidation(type='whole',operator='between',formula1='1',formula2='4',allow_blank=True,showErrorMessage=True,error='1~4'); wq.add_data_validation(dv); dv.add(f'M5:R{4+NI}')

# ---------- 보유·응답자정보
wp=wb.create_sheet('보유·응답자정보'); title(wp,'Part A 응답자·조직 정보(10) / Part B 프로세스 보유 확인(22)','Part B: 1 운영한다 · 2 부분적으로 운영한다 · 3 운영하지 않는다 · 4 모르겠다. 연결 문항은 출발·도착 프로세스가 모두 1·2일 때만 제시(survey/SURVEY_SPEC.md).')
hdr(wp,4,['ID','구분','항목','응답 방식','응답입력 열'],[8,10,36,52,12]); r=5
for i,a in enumerate(RESP['partA']):
    for j,v in enumerate([a['id'],'Part A',a['item'],a['options'],CL(2+i)]): put(wp,r,1+j,v)
    r+=1
for i,o in enumerate(OB):
    for j,v in enumerate([f'B{i+1:02d}','Part B',f"{o['id']} {o['name']}",'1 운영 / 2 부분 / 3 미운영 / 4 모름',CL(12+i)]): put(wp,r,1+j,v)
    r+=1

# ---------- 응답입력
wr=wb.create_sheet('응답입력'); title(wr,'응답 입력 (7행부터 최대 50명)','3행=연결 ID, 4행=품질 차원 코드(집계용 태그, 수정 금지). 6행은 예시(집계 제외). 연결·종합·검증 문항: 1~5 또는 9. 보유 확인: 1~4.')
FIRST,LAST=7,56
heads=['응답자ID']+[a['id'] for a in RESP['partA']]+[f'B{i:02d}' for i in range(1,23)]+[it['id'] for it in IT]
C1,C2=CL(resp_start),CL(resp_start+NI-1)
for j,h in enumerate(heads):
    c=wr.cell(5,1+j,h); c.font=F(True,'FFFFFF'); c.fill=HF; c.alignment=CT; c.border=BD; wr.column_dimensions[CL(1+j)].width=12 if j==0 else 8
for i,it in enumerate(IT):
    lt=it['link'] or ''
    dt=it['dim'] if it['part']=='C 연결' else ('G-'+it['dim'] if it['part']=='D 종합' and it['dim']!='ALL' else ('G-ALL' if it['part']=='D 종합' else it['id'][:2]))
    for rr,v in ((3,lt),(4,dt)):
        c=wr.cell(rr,resp_start+i,v); c.font=F(c='7F7F7F',s=8); c.alignment=CT
wr.cell(3,1,'연결 ID 태그').font=F(True,s=9); wr.cell(4,1,'차원 태그').font=F(True,s=9)
wr.cell(4,2,'Part A').font=F(True); wr.cell(4,12,'Part B').font=F(True)
ex={'A01':'제조업','A02':'300~999','A03':'3~5년','A04':'보유','A05':'BC 관리자','A06':'8','A07':'1','A08':'2','A09':'부분','A10':'일부'}
for j,h in enumerate(heads):
    c=wr.cell(6,1+j); v='EX-예시' if h=='응답자ID' else ex.get(h, 1 if h.startswith('B') and h[1:].isdigit() else [4,3,4,3,5,4,3,9,4][j%9])
    c.value=v; c.font=F(i=True,c='7F7F7F'); c.alignment=CT; c.fill=GREY; c.border=BD
for r in range(FIRST,LAST+1):
    for j in range(len(heads)):
        c=wr.cell(r,1+j); c.fill=INF; c.font=F(c=BLUE); c.alignment=CT; c.border=BD
dvL=DataValidation(type='list',formula1='"1,2,3,4,5,9"',allow_blank=True,showErrorMessage=True,error='1~5 또는 9'); dvB=DataValidation(type='list',formula1='"1,2,3,4"',allow_blank=True,showErrorMessage=True,error='1~4')
wr.add_data_validation(dvL); wr.add_data_validation(dvB)
dvB.add(f'L{FIRST}:{CL(33)}{LAST}'); dvL.add(f'{C1}{FIRST}:{C2}{LAST}'); wr.freeze_panes='B7'
def rng(i): return f'응답입력!${col_of[i]}${FIRST}:${col_of[i]}${LAST}'
LT=f'응답입력!${C1}$3:${C2}$3'; DT=f'응답입력!${C1}$4:${C2}$4'
def RR(r): return f'응답입력!${C1}{r}:${C2}{r}'

# ---------- 연결점수
wc=wb.create_sheet('연결점수'); title(wc,'응답자별 연결 점수 (적용된 품질 차원 문항의 평균, 9 제외, 최소 3문항 응답 시 산출)')
hdr(wc,5,['응답자ID']+[l['id'] for l in LK]+['연결 평균'],[12]+[7]*NL+[12]); lastc=CL(1+NL)
for i,r in enumerate(range(FIRST,LAST+1)):
    rr=6+i; put(wc,rr,1,f'=IF(응답입력!A{r}="","",응답입력!A{r})',al=CT)
    for j,l in enumerate(LK):
        put(wc,rr,2+j,f'=IF(COUNTIFS({LT},"{l["id"]}",{RR(r)},"<9")>=3,SUMIFS({RR(r)},{LT},"{l["id"]}",{RR(r)},"<9")/COUNTIFS({LT},"{l["id"]}",{RR(r)},"<9"),"")',al=CT,fmt='0.00')
    put(wc,rr,2+NL,f'=IFERROR(AVERAGE(B{rr}:{lastc}{rr}),"")',F(True),CT,fmt='0.00')
wc.freeze_panes='B6'

# ---------- 차원점수
wd=wb.create_sheet('차원점수'); title(wd,'응답자별 품질 차원 점수 (연결 문항 전체에서 차원별 평균, 9 제외) · 연결 품질 지수 = 6개 차원 평균(동일 가중치, 초기값)')
hdr(wd,5,['응답자ID']+[d['name'] for d in DM]+['연결 품질 지수'],[12]+[11]*6+[16])
for i,r in enumerate(range(FIRST,LAST+1)):
    rr=6+i; put(wd,rr,1,f'=IF(응답입력!A{r}="","",응답입력!A{r})',al=CT)
    for j,d in enumerate(DORD):
        mn=2 if d=='FC' else 5
        put(wd,rr,2+j,f'=IF(COUNTIFS({DT},"{d}",{LT},"L*",{RR(r)},"<9")>={mn},SUMIFS({RR(r)},{DT},"{d}",{LT},"L*",{RR(r)},"<9")/COUNTIFS({DT},"{d}",{LT},"L*",{RR(r)},"<9"),"")',al=CT,fmt='0.00')
    put(wd,rr,8,f'=IFERROR(AVERAGE(B{rr}:G{rr}),"")',F(True),CT,fmt='0.00')
wd.freeze_panes='B6'

# ---------- 점수요약
w2=wb.create_sheet('점수요약'); title(w2,'점수 요약 — 연결·차원·영역·증거유형·관계동사별 평균 (응답입력 7~56행, 9 제외)')
hdr(w2,4,['ID','연결(S→T)','영역','증거유형','관계동사']+[d['name'] for d in DM]+['연결 점수','응답 수','가장 낮은 차원'],[7,34,22,9,16,10,10,10,10,10,10,11,9,16])
for i,l in enumerate(LK):
    r=5+i
    put(w2,r,1,l['id']); put(w2,r,2,f"{l['source']} → {l['target']}"); put(w2,r,3,l['domain']); put(w2,r,4,l['evidence'],al=CT); put(w2,r,5,VBD[l['verb']]['verb'])
    for j,d in enumerate(DORD):
        if d in l['dims']: put(w2,r,6+j,f'=IFERROR(AVERAGEIF({rng(l["id"]+"-"+d)},"<9"),"")',al=CT,fmt='0.00')
        else: put(w2,r,6+j,'-',al=CT,font=F(c='98A2B3'))
    put(w2,r,12,f'=IFERROR(AVERAGE(F{r}:K{r}),"")',F(True),CT,fmt='0.00')
    put(w2,r,13,'='+'+'.join(f'COUNTIF({rng(l["id"]+"-"+d)},"<9")' for d in l['dims']),al=CT)
    put(w2,r,14,f'=IF(L{r}="","",INDEX($F$4:$K$4,MATCH(MIN(F{r}:K{r}),F{r}:K{r},0)))',al=CT)
e=4+NL; r0=e+2
put(w2,r0,1,'차원별 평균',F(True))
for j,d in enumerate(DM): put(w2,r0+1+j,2,d['name']); put(w2,r0+1+j,3,f'=IFERROR(AVERAGE({CL(6+j)}5:{CL(6+j)}{e}),"")',F(True),CT,fmt='0.00')
r1=r0+8; put(w2,r1,1,'영역별 평균',F(True))
for i,d in enumerate(DOMAIN_ORDER): put(w2,r1+1+i,2,d); put(w2,r1+1+i,3,f'=IFERROR(AVERAGEIF($C$5:$C${e},B{r1+1+i},$L$5:$L${e}),"")',F(True),CT,fmt='0.00')
r2=r1+7; put(w2,r2,1,'증거유형별 평균',F(True))
for i,t in enumerate(['E1','E2','E3']): put(w2,r2+1+i,2,t); put(w2,r2+1+i,3,f'=IFERROR(AVERAGEIF($D$5:$D${e},B{r2+1+i},$L$5:$L${e}),"")',F(True),CT,fmt='0.00')
r3=r2+5; put(w2,r3,1,'관계동사별 평균',F(True)); used=[v for v in VB if v['links']]
for i,v in enumerate(used): put(w2,r3+1+i,2,v['verb']); put(w2,r3+1+i,3,f'=IFERROR(AVERAGEIF($E$5:$E${e},B{r3+1+i},$L$5:$L${e}),"")',F(True),CT,fmt='0.00')
r4=r3+len(used)+2; put(w2,r4,2,f'연결 평균(연결 {NL}개)',F(True)); put(w2,r4,3,f'=IFERROR(AVERAGE(L5:L{e}),"")',F(True),CT,fmt='0.00')
w2.freeze_panes='C5'

# ---------- 검증
wv=wb.create_sheet('검증'); title(wv,'검증 — 기준값 · 구조 점검 · 자료 기반 점검 · 검증 계획','기준값(파란 글씨)은 일반적 관례이며 연구자가 조정')
hdr(wv,3,['기준','값','설명'],[44,16,80])
for i,(a,b,c) in enumerate([('내용타당도 I-CVI 기준',0.78,'전문가 6~10인 패널의 관례(정의서 D23: 연구자 선택, 논문 근거 없음)'),('Arias 합의 채택 기준',0.8,'정의서 D22 / 인용 A43·A44'),('Arias 합의 기각 기준',0.2,'< 이 값이면 삭제'),('Cronbach α / ω 기준',0.7,'차원별 문항 신뢰도(통계 프로그램에서 산출)'),('문항-총점 상관 기준',0.3,'통계 프로그램'),('차원 간 상관 상한',0.85,'6개 차원이 서로 구분되는지(변별타당도). 넘으면 차원 통합 검토')]):
    put(wv,4+i,1,a); put(wv,4+i,2,b,F(c=BLUE),CT,INF); put(wv,4+i,3,c)
r=11; hdr(wv,r,['구조 점검','현재 값','기대 값','결과'])
nA=sum(len(l['dims']) for l in LK)
chk=[('총 문항 수',f'=COUNTA(문항!A5:A{4+NI})',NI),('연결 문항 수(Part C)',f'=COUNTIF(문항!B5:B{4+NI},"C 연결")',nA)]
for d in DM: chk.append((f'{d["name"]} 문항 수',f'=COUNTIF(문항!G5:G{4+NI},"{d["name"]}")',sum(1 for l in LK if d['id'] in l['dims'])))
chk+=[('기대 문항 수와 일치하는 연결 수',f'=COUNTIF(연결!U{LR1}:U{LR2},"OK")',NL),('적용매트릭스 합계',f'=적용매트릭스!H{5+NL}',nA),
('증거유형 E1 연결 수',f'=COUNTIF(연결!H{LR1}:H{LR2},"E1")',sum(1 for l in LK if l['evidence']=='E1')),('증거유형 E2 연결 수',f'=COUNTIF(연결!H{LR1}:H{LR2},"E2")',sum(1 for l in LK if l['evidence']=='E2')),('증거유형 E3 연결 수',f'=COUNTIF(연결!H{LR1}:H{LR2},"E3")',sum(1 for l in LK if l['evidence']=='E3')),
('환류 연결 수',f'=COUNTIF(연결!J{LR1}:J{LR2},"Y")',8),('집단 노드 연결 수',f'=COUNTIF(연결!I{LR1}:I{LR2},"집단")',2),
('관계동사 사용 수(사용 안 한 V10 제외)',f'=COUNTIF(관계동사!J5:J{4+len(VB)},">0")',len(used)),('객체(프로세스) 수','=COUNTA(객체!A5:A26)',22),
('인용문 수',f'=COUNTA(인용근거!A5:A{4+NQ})',NQ),('원문 대조 일치 인용 수',f'=COUNTIF(인용근거!H5:H{4+NQ},"일치*")',NQ),('외부 학술 근거 후보 수',f'=COUNTA(학술근거!A5:A{4+NS})',NS)]
for i,(a,f_,ex_) in enumerate(chk):
    rr=r+1+i; put(wv,rr,1,a); put(wv,rr,2,f_,al=CT); put(wv,rr,3,ex_,F(c=BLUE),CT); put(wv,rr,4,f'=IF(B{rr}=C{rr},"OK","확인")',F(True),CT)
rr=r+len(chk)+2; put(wv,rr,1,'외부 문헌 원문 확인 현황 (학술근거 시트 "확인 완료" 열)',F(True),fill=GREY)
put(wv,rr+1,1,'확인 완료(Y) 건수'); put(wv,rr+1,2,f'=COUNTIF(학술근거!L5:L{4+NS},"Y")',al=CT)
put(wv,rr+2,1,'미확인 건수'); put(wv,rr+2,2,f'={NS}-B{rr+1}',al=CT); put(wv,rr+2,3,'연구자가 원문 확인 전에는 외부 문헌을 인용 근거로 확정하지 말 것')
rr+=4; put(wv,rr,1,'전문가 판정(문항 시트)',F(True),fill=GREY)
put(wv,rr+1,1,'채택 문항 수'); put(wv,rr+1,2,f'=COUNTIF(문항!T5:T{4+NI},"채택")',al=CT)
put(wv,rr+2,1,'수정·토론 문항 수'); put(wv,rr+2,2,f'=COUNTIF(문항!T5:T{4+NI},"수정·토론")',al=CT)
put(wv,rr+3,1,'평가 미입력 문항 수'); put(wv,rr+3,2,f'={NI}-B{rr+1}-B{rr+2}',al=CT)
rr+=5; hdr(wv,rr,['자료 기반 점검(응답 입력 후)','값','기대','결과'])
pairs=[('역문항 일관성: DQ01(정합성) vs RV1','DQ01','RV1'),('역문항 일관성: DQ02(추적성) vs RV2','DQ02','RV2'),('역문항 일관성: DQ05(최신성) vs RV3','DQ05','RV3'),('역문항 일관성: DQ06(환류폐쇄성) vs RV4','DQ06','RV4')]
for i,(a,x,y) in enumerate(pairs):
    q=rr+1+i; put(wv,q,1,a); put(wv,q,2,f'=IFERROR(CORREL({rng(x)},{rng(y)}),"자료 부족")',al=CT,fmt='0.00'); put(wv,q,3,'음(−)의 상관'); put(wv,q,4,f'=IF(ISNUMBER(B{q}),IF(B{q}<0,"OK","확인"),"")',F(True),CT)
q=rr+1+len(pairs)
for j,d in enumerate(DM):
    put(wv,q,1,f'수렴: {d["name"]} 점수 vs 종합 문항 DQ0{j+1}'); put(wv,q,2,f'=IFERROR(CORREL(차원점수!{CL(2+j)}6:{CL(2+j)}55,{rng("DQ0"+str(j+1))}),"자료 부족")',al=CT,fmt='0.00'); put(wv,q,3,'양(+)의 상관'); put(wv,q,4,f'=IF(ISNUMBER(B{q}),IF(B{q}>0,"OK","확인"),"")',F(True),CT); q+=1
put(wv,q,1,'준거: 연결 품질 지수 vs CR1'); put(wv,q,2,f'=IFERROR(CORREL(차원점수!H6:H55,{rng("CR1")}),"자료 부족")',al=CT,fmt='0.00'); put(wv,q,3,'양(+)의 상관'); put(wv,q,4,f'=IF(ISNUMBER(B{q}),IF(B{q}>0,"OK","확인"),"")',F(True),CT); q+=1
put(wv,q,1,'수렴: 연결 품질 지수 vs 전반 문항 DQ07'); put(wv,q,2,f'=IFERROR(CORREL(차원점수!H6:H55,{rng("DQ07")}),"자료 부족")',al=CT,fmt='0.00'); put(wv,q,3,'양(+)의 상관'); put(wv,q,4,f'=IF(ISNUMBER(B{q}),IF(B{q}>0,"OK","확인"),"")',F(True),CT); q+=2
put(wv,q,1,'변별: 차원 간 상관 행렬 (차원점수 시트 기준)',F(True),fill=GREY); q+=1
put(wv,q,1,'',al=CT)
for j,d in enumerate(DM): c=wv.cell(q,2+j,d['name']); c.font=F(True,'FFFFFF'); c.fill=HF; c.alignment=CT; c.border=BD
mat0=q+1; offd=[]
for i,di in enumerate(DM):
    put(wv,mat0+i,1,di['name'],F(True))
    for j,dj in enumerate(DM):
        put(wv,mat0+i,2+j,f'=IFERROR(CORREL(차원점수!{CL(2+i)}6:{CL(2+i)}55,차원점수!{CL(2+j)}6:{CL(2+j)}55),"")',al=CT,fmt='0.00')
        if j>i: offd.append(f'{CL(2+j)}{mat0+i}')
for j in range(6): wv.column_dimensions[CL(2+j)].width=max(wv.column_dimensions[CL(2+j)].width or 0,14)
q=mat0+7; put(wv,q,1,'차원 간 최대 상관'); put(wv,q,2,'=IFERROR(MAX('+','.join(offd)+'),"자료 부족")',al=CT,fmt='0.00'); put(wv,q,3,'< 차원 간 상관 상한(B9)'); put(wv,q,4,f'=IF(ISNUMBER(B{q}),IF(B{q}<$B$9,"OK","확인(차원 통합 검토)"),"")',F(True),CT)
q+=2; put(wv,q,1,'참고: CORREL은 9(해당없음)도 포함한다. 실제 분석은 9를 결측 처리한 자료로 통계 프로그램에서 다시 산출한다. 역문항(RV1~4, CR4)은 분석 시 6−값으로 뒤집어 사용한다. 환류폐쇄성 점수는 환류 연결 8개 문항만으로 계산되므로 다른 차원과 문항 수가 다르다.',F(i=True,c='475467'))
wv.merge_cells(start_row=q,start_column=1,end_row=q,end_column=4); wv.row_dimensions[q].height=50
q+=2; hdr(wv,q,['검증 단계','방법','기준 / 산출 위치'])
plan=[('1 인용 검증','Arias·Białas 인용을 PDF 원문과 자동 대조(scripts/verify_quotes.py)',f'인용 {NQ}/{NQ} 일치(완료)'),
('2 외부 문헌 확인','연구자가 외부 문헌 원문에서 정의·문장을 찾아 학술근거 시트에 기입','학술근거 시트 "확인 완료" 열 — 미확인 문헌은 근거로 확정 금지'),
('3 구조 점검','추적성(문항=연결 1×차원 1), 적용매트릭스와 문항 수 일치','위 구조 점검(자동)'),
('4 전문가 내용타당도','6~10인: 연결 타당도(E2·E3 우선), 문항 적합도, 차원 귀속 맹검 재분류','I-CVI ≥ 0.78 / 일치율 ≥ 80% 채택 · 20~80% 토론 · < 20% 삭제(Arias A44 준용)'),
('5 인지면담','5~8인 설명시키기(출발·도착, 6개 차원 구분)','수기'),
('6 신뢰도·구조','차원별 α/ω, 문항-총점 상관, 6차원 CFA(형성적 2차 요인 포함)','통계 프로그램(150명 이상 권장)'),
('7 변별·수렴','차원 간 상관 < 0.85, 차원 점수와 종합 문항 DQ 상관, 역문항 일관성','본 시트'),
('8 준거·집단','연결 품질 지수 vs CR1~4, 증거유형별(E1/E2/E3) 비교, ISO 22301 인증(A04) 집단 비교','점수요약 / 통계 프로그램'),
('9 활용 검증','점수로 "어느 연결의 어느 차원이 약한가" 진단이 되는지(Białas B08 방식)','점수요약')]
for i,(a,b,c) in enumerate(plan):
    put(wv,q+1+i,1,a); put(wv,q+1+i,2,b); put(wv,q+1+i,3,c); wv.row_dimensions[q+1+i].height=36
wv.column_dimensions['B'].width=50

order=['안내','정의서','품질차원','학술근거','인용근거','연결','적용매트릭스','관계동사','보류연결','객체','트리플','관계매트릭스','온톨로지맵','문항','보유·응답자정보','응답입력','연결점수','차원점수','점수요약','검증']
wb._sheets=[wb[n] for n in order]
for s in wb.worksheets: s.sheet_view.showGridLines=s.title not in('안내','온톨로지맵')
wb.save(out('BCMS_연계성_V3.xlsx')); print('xlsx saved')

# ---- survey json/csv
sv=[]; objs_by={o['id']:f'B{i+1:02d}' for i,o in enumerate(OB)}
for it in IT:
    lk=LKD.get(it['link'])
    sv.append(dict(id=it['id'],part=it['part'],link_id=it['link'] or None,quality_dimension=it['dim'] or None,dimension_name=it['state'],text=it['text'],reverse_scored=it['reverse'],
        scale='likert5_na9',relation_verb=VBD[lk['verb']]['verb'] if lk else None,evidence_type=lk['evidence'] if lk else None,domain=lk['domain'] if lk else None,
        requires_process_ids=[objs_by[x] for x in (lk['source'],lk['target']) if x in objs_by] if lk else [],basis_type=it['basis_type'],basis_quotes=it['basis_quotes']))
meta=dict(version='V3',construct='연결 품질(6개 차원)',dimensions=[dict(id=d['id'],name=d['name'],definition=d['researcher_def']) for d in DM],
          scale=dict(type='likert5_na9',labels={'1':'전혀 그렇지 않다','2':'그렇지 않다','3':'보통이다','4':'그렇다','5':'매우 그렇다','9':'해당 없음/모르겠음(결측)'}),
          process_presence=dict(scale={'1':'운영한다','2':'부분적으로 운영한다','3':'운영하지 않는다','4':'모르겠다'},items=[dict(id=f'B{i+1:02d}',process_id=o['id'],name=o['name']) for i,o in enumerate(OB)]),
          respondent_profile=RESP['partA'],items=sv)
os.makedirs(os.path.join(ROOT,'survey'),exist_ok=True)
json.dump(meta,open(os.path.join(ROOT,'survey','survey_items_v3.json'),'w',encoding='utf-8'),ensure_ascii=False,indent=1)
with open(os.path.join(ROOT,'survey','survey_items_v3.csv'),'w',encoding='utf-8-sig',newline='') as f:
    w_=csv.writer(f); w_.writerow(['id','part','link_id','quality_dimension','dimension_name','text','reverse_scored','scale','relation_verb','evidence_type','domain','requires_process_ids','basis_type','basis_quotes'])
    for s in sv: w_.writerow([s['id'],s['part'],s['link_id'] or '',s['quality_dimension'] or '',s['dimension_name'],s['text'],int(s['reverse_scored']),s['scale'],s['relation_verb'] or '',s['evidence_type'] or '',s['domain'] or '',';'.join(s['requires_process_ids']),s['basis_type'],';'.join(s['basis_quotes'])])
print('survey saved')

# ---- md
L=[]; A=L.append
A('# BCMS 프로세스 연결 품질 — 근거·정의서 (V3)\n')
A('> 박사논문 프로젝트 · 근거 문헌: Arias-Aranda et al.(2026) *Appl. Sci.* 16, 3219 / Białas(2010) *Ontological approach to the business continuity management system development*\n')
A(f'> Arias·Białas의 영문 인용문은 PDF 원문과 자동 대조했다({ver["quotes_ok"]}/{ver["quotes_total"]} 일치, 검증일 {VDATE}). 외부 학술 문헌 {NS}건은 서지만 웹 검색으로 확인했고 원문 문장은 확인하지 못했다(원문 대조 {nvsrc}건). 한국어 번역은 연구자 번역(참고용)이다.\n')
A('## 0. 읽는 법\n\n- **▶ 인용**: 논문 원문을 그대로 옮긴 부분(쪽 번호 포함, 원문 대조 완료).\n- **▶ 연구자 의견**: 인용에 없는 해석·가정·분류·정의. 박사논문에서 연구자가 직접 확정해야 한다.\n- **연결의 증거유형** E1=원문이 S→T 연결을 직접 서술 / E2=출발·도착 일부를 연구자가 보완 / E3=원문에 서술 없음(연구자 가정).\n- **품질 차원의 근거 등급** T1=Arias·Białas 원문 인용 / T2=외부 문헌(서지만 확인) / T3=연구자 정의.\n')
A('## 1. 3층 모형\n\n| 층 | 질문 | 내용 | 근거 |\n|---|---|---|---|\n| 객체층 | 무엇과 무엇을 연결하는가 | 프로세스 22개와 그 산출물 | Arias Table 5 (인용) |\n| 관계층 | 어떤 방식으로 연결되는가 | 관계동사 11개 (Arias 서술문에서 추출) | Arias §4.3 (인용) + 동사 분류(연구자) |\n| 품질층 | 그 연결이 얼마나 잘 작동하는가 | 6개 차원 | 첨부 문서의 초기 가설(연구자) + T1/T2 근거 |\n')
A('문항 = **객체쌍 + 관계동사 + 품질조건**. 예) "요구사항 목록이(객체 S) BIA의 우선순위·복구목표 설정에(객체 T) 입력될 때(관계동사: 입력된다), 두 내용이 서로 모순 없이 일치한다(품질조건: 정합성)."\n')
A('## 2. 품질 6개 차원과 근거\n')
for d in DM:
    A(f"### {d['order']}. {d['name']} ({d['en']})\n")
    A(f"- 첨부 문서의 정의(연구자 제시): {d['attachment_def']}\n- ▶ 연구자 의견(본 연구의 조작적 정의): {d['researcher_def']}\n- 문항 틀: `{d['item_form']}`\n- **T1 근거 강도: {d['t1_strength']}** — {d['t1_note']}")
    for k in d['t1']+d.get('t1_weak',[]):
        q=Q[k]; A(f"  - ▶ 인용 [{k}] ({q['source']} p.{q['pages']}, {q['section']}): \"{q['text']}\"\n    - 번역: {q['ko']}")
    A(f"- T2 외부 문헌 후보: {', '.join(d['ext'])} — {d['ext_note']}\n- ▶ 연구자 의견(다른 차원과의 겹침): {d['overlap']}\n")
A('## 3. 외부 학술 근거 후보 (서지 확인, 원문 미확인)\n\n네트워크 정책으로 원문을 열람하지 못했다. 아래 "요지"는 검색 결과 요약이며 인용이 아니다. 연구자가 원문에서 정의·문장을 확인한 뒤 근거로 확정해야 한다.\n')
for s in SRC:
    dm=', '.join(DMD[d]['name'] for d in s['dims']) if s['dims'] else '방법(구성·검증)'
    A(f"- **{s['id']}** [{s['type']}] {s['citation']}"+(f" doi:{s['doi']}" if s['doi'] else '')+f"\n  - 관련 차원: {dm}\n  - 요지(인용 아님): {s['claim']}\n  - 상태: {s['bib_status']} / {s['text_status']}")
A('\n## 4. 방법 근거 (Białas 인용)\n')
for k in ['B02','B03','B04','B05','B06','B07','B08','B14','B15']:
    q=Q[k]; A(f"- **{k}** (Białas p.{q['pages']}, {q['section']}) ▶ 인용: \"{q['text']}\"\n  - 번역: {q['ko']}\n  - 사용: {q['use']}")
A('\n## 5. 정의서\n\n| ID | 용어 | 유형 | 정의 | 근거 인용 | 근거·한계 |\n|---|---|---|---|---|---|')
for d in DF: A(f"| {d['id']} | {d['term']} | {d['type']} | {d['definition']} | {', '.join(d['quotes']) or '—'} | {d['note'] or '—'} |")
A('\n## 6. 관계동사 11개\n\n| ID | 관계동사 | 의미(S→T) | Arias 원문 동사 | 근거 인용 | 첨부 문서와의 대응 | 사용 연결 |\n|---|---|---|---|---|---|---|')
for v in VB: A(f"| {v['id']} | {v['verb']} | {v['definition']} | {v['arias_verbs']} | {', '.join(v['quotes'])} | {v['attachment_map']} | {', '.join(v['links']) or '미사용'} |")
A(f'\n## 7. 연결 {NL}개 — 인용과 연구자 의견\n')
for l in LK:
    A(f"### {l['id']} · {nm(l['source'])} → {nm(l['target'])}\n")
    A(f"- 영역: {l['domain']} · 관계동사: **{VBD[l['verb']]['verb']}** · **증거유형 {l['evidence']}**{' · 집단 노드 포함' if l['group'] else ''} · 적용 차원: {', '.join(DMD[d]['name'] for d in l['dims'])}")
    if l['quotes_primary']:
        for k in l['quotes_primary']:
            q=Q[k]; A(f"- ▶ 인용 [{k}] (Arias p.{q['pages']}, {q['section']}): \"{q['text']}\"\n  - 번역: {q['ko']}")
    else: A('- ▶ 인용: 원문에 S→T 연결 서술 없음')
    for k in l['quotes_support']:
        q=Q[k]; A(f"- ▷ 보조 인용 [{k}] ({q['source']} p.{q['pages']}, {q['section']}): \"{q['text']}\"\n  - 번역: {q['ko']}")
    A(f"- 인용이 확립하는 사실: {l['established']}\n- ▶ 연구자 의견(해석·보완·가정): {l['researcher']}\n- ▶ 연구자 의견(차원 적용): {l['quality_note']}\n- Białas 대응: {l['bialas']}\n- 검증 필요: {l['validation']}\n")
    for i in IT:
        if i['link']==l['id']: A(f"  - {i['id']} ({i['state']}): {i['text']}")
    A('')
A('## 8. 보류·제외한 후보 연결\n\n| ID | 출발→도착 | 후보 | 채택하지 않은 이유 | 권고 |\n|---|---|---|---|---|')
for h in HL: A(f"| {h['id']} | {h['source']}→{h['target']} | {h['idea']} | {h['reason']} | {h['decision']} |")
A('\n## 9. 종합·검증 문항\n')
for i in IT:
    if i['part'][0] in 'DE': A(f"- {i['id']} ({i['state']}{', 역문항' if i['reverse'] else ''}): {i['text']}"+(f"  \n  근거: {i['basis_type']}"+(f" · 인용 {', '.join(i['basis_quotes'])}" if i['basis_quotes'] else '')))
A('\n## 10. 한계\n\n1. 연결 31개는 전문가 검증 전 가설이다. Arias는 프로세스가 핵심인지를 검증했을 뿐 연결을 검증하지 않았다(인용 A06).\n2. 6개 차원과 "연결 품질" 용어는 연구자가 제시한 초기 가설이다. 두 논문 안에서 추적성은 근거가 없고, 정합성·증빙성·책임성·최신성은 간접 근거, 환류폐쇄성은 비교적 직접 근거다.\n3. 외부 학술 문헌 12건은 서지만 확인했고 원문은 열람하지 못했다. 원문 확인 전에는 인용 근거로 확정할 수 없다.\n4. 6개 차원은 서로 겹칠 수 있어(특히 정합성·추적성·증빙성) 변별타당도 검증이 필요하다. 환류폐쇄성은 환류 연결 8개에만 적용되어 다른 차원과 문항 수가 다르다.\n5. Białas(2010)는 BS 25999 기반 프로토타입이며(B14, B15) 본 연구는 방법만 차용했다.\n6. 통계 기준(I-CVI 0.78, α 0.70 등)은 관례이며 논문 근거가 없다.\n')
open(out('BCMS_연계성_V3_근거정의서.md'),'w',encoding='utf-8').write('\n'.join(L)); print('md saved')
