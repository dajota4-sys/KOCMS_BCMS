"""V5 엑셀: V4 연구자 워크북(기준) + 이야기하기·인용지도·참고문헌·관계MAP이미지 시트. data/v5/*.json -> outputs/BCMS_연계성_V5.xlsx"""
import json,os,sys,shutil
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image as XImg
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC=sys.argv[1] if len(sys.argv)>1 else os.path.join(ROOT,'data/v5/v4_workbook.xlsx')
D=json.load(open(os.path.join(ROOT,'data/v5/v4_baseline.json'),encoding='utf-8'))
ST=json.load(open(os.path.join(ROOT,'data/v5/story.json'),encoding='utf-8'))
wb=openpyxl.load_workbook(SRC)
FN='Noto Sans KR'
HF=Font(name=FN,sz=11,b=True,color='FFFFFF'); HFILL=PatternFill('solid',fgColor='17324D')
BF=Font(name=FN,sz=11); thin=Side(style='thin',color='D5DAE3'); BR=Border(left=thin,right=thin,top=thin,bottom=thin)
def sheet(name,head,rows,widths,tab=None,idx=None):
    ws=wb.create_sheet(name,idx); 
    if tab: ws.sheet_properties.tabColor=tab
    ws.append(head)
    for c in ws[1]: c.font=HF; c.fill=HFILL; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); c.border=BR
    ws.row_dimensions[1].height=32
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=BF; c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=BR
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    ws.freeze_panes='A2'; return ws
# 1) 시작점 갱신
ws=wb['시작점']
ws['A1']='상태'; ws['B1']='연구 기준 2026년 10월 9일 V5 (V4를 시작점으로 보존한 발표용 최종본)'
ws['A16']='읽는 순서'; ws['B16']='시작점 → 이야기하기 → 관계MAP(또는 관계MAP이미지) → 관계근거31 → 예비문항179. 인용지도·참고문헌·나머지 시트는 근거와 수정 이력을 확인할 때 사용.'
ws.append(['V5 변경 범위','V4의 연구 내용(분류·문항·정의·Gate)은 변경하지 않았다. V5는 발표 순서(이야기하기), 개념별 인용 위치(인용지도), 참고문헌과 확인 수준, 22×22 MAP 이미지를 추가했다.'])
ws.append(['인용 확인 수준','① 이 세션의 PDF 문자열 대조(Arias·Białas 72개) ② 연구자 확인(V4 기록 — 이 세션에서 재확인하지 못함) ③ 연구자 해석·정의. 시트 참고문헌 참조.'])
for r in (17,18):
    for c in ('A','B'):
        cell=ws[f'{c}{r}']; cell.font=Font(name=FN,sz=11,b=(c=='A')); cell.alignment=Alignment(wrap_text=True,vertical='top')
    ws.row_dimensions[r].height=48
# 2) 이야기하기
rows=[[c['no'],c['title'],c['slides'],c['message'],c['sheets'],c['cites']] for c in ST['chapters']]
sheet('이야기하기',['장','이야기','발표 슬라이드','핵심 메시지','근거 시트','대표 인용·출처'],rows,[6,18,12,70,26,44],tab='FF2E9E6B',idx=1)
# 3) 인용지도: 개념 -> 학술근거 -> 위치 -> 확인수준 -> 한계
qc={r['코드']:r for r in D['품질개념6']}
ref={r['ID']:r for r in D['외부학술근거']}
LV={'S01':'연구자 본문 확인','S03':'연구자 본문 확인','S04':'연구자 본문 확인','S06':'초록만 확인','S07':'초록만 확인','S10':'연구자 본문 확인','S13':'연구자 본문 확인','S14':'초록만 확인','S15':'연구자 본문 확인'}
cm=[('CO','S03','Appendix D1','표현 일관성은 형식·표현 호환성. 의미적 정합성의 정의로 사용하지 않음'),
('CO','S06','저자기관 초록; 저널 pp.993–1009','요구사항 일관성·진화·검증 논의. 본문 직접 보강 필요'),
('CO','A12·A14·A16','Arias PDF p.15–16 §4.3','BIA→RA, BIA→STR, STR→PLAN 연결 서술(관계 근거). 정합성 정의의 근거는 아님'),
('TR','S01','§5.1 PDF p.4; 저널 pp.94–101','전방·후방 추적. 현재 문항은 후방 근거 추적만'),
('EV','B17','Białas PDF p.5','훈련 결과·기록의 증거 역할. 연결 전체의 품질 차원으로 제시되지는 않음'),
('EV','A13·A37','Arias','연구자 확장. TR과 중복 검토 필요'),
('AC','A49·A50','Arias §4.3','프로세스 전체 책임 서술'),
('AC','S07','출판사 초록','Accountability = 설명·정당화 의무와 질문·판단·결과 부담. 담당 지정(AC)과 다름'),
('TM','S04','PDF pp.4–5; 저널 pp.214–215','과업 맥락의 적시성과 currency 구분. 변경 전달 기한은 연구자 정의'),
('TM','A27·A28·A41','Arias','변경관리 서술(관계 근거)'),
('FC','A23·A25·A29','Arias §4.3 pp.15–16','평가 결과의 개선 활용 서술. 조치 완료까지의 닫힌 순환은 실증하지 않음 → 개념 보류'),
('FC','S08','직접 근거에서 제외','Double-loop는 지배 규칙·가정 재검토. 개선조치 종결과 등치하지 않음'),
('측정모형','S10','저널 pp.302–303; PDF pp.10–11','형성적·반영적 선택은 개념 정의에 따름. 6차원 평균을 보증하지 않음'),
('척도개발','S13','Table 1','영역·문항 생성·내용타당화·인지면접 절차. 고정 문항 수 보증 아님'),
('내용타당도','S14','출판사 초록','I-CVI .78 이상 권고 + 우연 합의 보정. 6인 5/6=.833 최소 충족'),
('온톨로지 범위','S15','PDF pp.4–5 §3 Step 1','범위·용도·질의로 영역 설정. 미완성 모델의 타당성을 보증하지 않음'),
('온톨로지 표현','Białas','§2 pp.2–3','클래스·인스턴스·object/data slot·역량질의(방법만 차용)')]
rows=[]
for k,s,loc,lim in cm:
    nm=qc[k]['수정 개념'] if k in qc else k
    role=qc[k]['검토 역할'] if k in qc else '방법·절차'
    lv=LV.get(s,'제공 PDF 문자열 대조(이 세션)' if (s.startswith('A') or s=='Białas' or s=='B17') else '')
    rows.append([k,nm,role,s,loc,lv,lim])
sheet('인용지도',['코드','개념','역할','근거','위치','확인 수준','범위·한계(이 근거가 말하지 않는 것)'],rows,[10,18,16,16,28,22,70],tab='FFE8A317',idx=5)
# 4) 참고문헌
rows=[[r['key'],r['full'],r['doi'],r['level'],r['use']] for r in ST['refs']]
sheet('참고문헌',['본문 인용 표기','참고문헌','DOI','확인 수준','본 연구에서의 용도'],rows,[28,80,26,36,44],tab='FFE8A317',idx=6)
BQ=json.load(open(os.path.join(ROOT,'data/v5/background_quotes.json'),encoding='utf-8'))
sheet('배경인용18',['인용ID','출처','PDF 쪽','절','원문','참고 번역(연구자)','사용처','문자 대조'],[[q['id'],q['source'],q['pages'],q['section'],q['text'],q['ko'],q['use'],'문자 대조 일치 2026-10-09'] for q in BQ],[8,10,8,10,70,60,34,22],tab='FFE8A317',idx=wb.sheetnames.index('인용원문72')+1)
# 5) 이미지
ws=wb.create_sheet('관계MAP이미지',3); ws.sheet_properties.tabColor='FF245E94'
img=XImg(os.path.join(ROOT,'outputs/BCMS_관계MAP_22x22_V5.png')); img.width=img.width*0.55; img.height=img.height*0.55; ws.add_image(img,'A1')
# 5b) V4 기록의 쪽 번호 오기 정정(B17: PDF p.6 -> p.5)
fixed=0
for w in wb:
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value,str) and 'B17 PDF p.6' in c.value: c.value=c.value.replace('B17 PDF p.6','B17 PDF p.5'); fixed+=1
print('B17 쪽 정정',fixed)
# 6) 수정이력 V5
ws=wb['수정이력']
add=[('R21','안내','V5 시작점','V4가 발표 순서·인용 위치를 한눈에 보여 주지 않음','이야기하기·인용지도·참고문헌 시트를 추가. 연구 분류·문항·정의는 V4와 동일.','본 작업'),
('R22','중요','인용 확인 수준 표시','Arias·Białas 문자 대조와 외부 문헌 확인 수준이 한 줄에 섞일 위험','① 이 세션 PDF 문자열 대조 ② 연구자 확인(V4 기록, 재확인 못 함) ③ 연구자 해석으로 3단계 구분.','참고문헌·인용지도'),
('R23','보완','72개 인용 재대조','V5 작성 시점의 일치 재확인','Arias·Białas 72개: 단일쪽 70개 문자 일치, 쪽 경계 A06(3–4)·A51(16–17)은 쪽 사이 머리글 때문에 간이 대조에서 불일치하나 PDF 직접 대조(전 72개 일치, 2026-10-09) 결과 유지. 해석·번역의 타당성 증거 아님.','outputs/quote_verification.json'),
('R24','보완','서지 표기','Boateng 외 저자 전체·Arias 권호 등 불완전','Boateng 외(2018) 저자 5인 표기. Arias는 권(16)·논문번호(3219)만 표기. Białas 게재지·연도·DOI는 미확정으로 남김.','참고문헌'),
('R26','중요','B17 쪽 번호','V4 개념 근거·수정이력에 B17을 PDF p.6으로 표기(인용원문72는 p.5)','Białas 제공 PDF 5쪽에서 B17 원문 확인(6쪽에는 없음). 예비문항179·품질개념6·수정이력의 \'B17 PDF p.6\'을 p.5로 정정.','Białas PDF 쪽 대조'),
('R25','보완','22×22 MAP 이미지','엑셀 격자 외에 발표용 이미지 필요','관계근거31에서 자동 생성한 PNG를 시트와 PPT에 삽입. 셀 29 + 집단 주석 2.','scripts/mkmatrix.py')]
add.sort()
for r in add:
    ws.append(list(r))
    for c in ws[ws.max_row]: c.font=BF; c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=BR
os.makedirs(os.path.join(ROOT,'outputs'),exist_ok=True)
wb.save(os.path.join(ROOT,'outputs/BCMS_연계성_V5.xlsx')); print(wb.sheetnames)
