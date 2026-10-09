"""V5 Word 해설서: data/v5/*.json + data/quotes.json -> outputs/BCMS_연계성_해설서_V5.docx"""
import json,os
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
J=lambda p:json.load(open(os.path.join(ROOT,p),encoding='utf-8'))
B=J('data/v5/v4_baseline.json');ST=J('data/v5/story.json');Q={q['id']:q for q in J('data/quotes.json')}
LK=B['관계근거31'];IT={r['문항ID']:r for r in B['예비문항179']};CQ={r['코드']:r for r in B['품질개념6']}
NAVY=RGBColor(0x14,0x21,0x3D);BLUE=RGBColor(0x2F,0x6D,0xB5);ORG=RGBColor(0xC8,0x6A,0x10);GRY=RGBColor(0x6B,0x7A,0x90)
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for m in('left_margin','right_margin'):setattr(sec,m,Cm(2.2))
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.0)
FN='맑은 고딕'
def setfont(st,sz,bold=None,color=None):
    st.font.name=FN;st.font.size=Pt(sz);st.element.rPr.rFonts.set(qn('w:eastAsia'),FN)
    if bold is not None:st.font.bold=bold
    if color:st.font.color.rgb=color
setfont(d.styles['Normal'],10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.3
for n,sz in(('Heading 1',18),('Heading 2',13.5),('Heading 3',11.5)):setfont(d.styles[n],sz,True,NAVY)
d.styles['Heading 1'].paragraph_format.space_before=Pt(18);d.styles['Heading 1'].paragraph_format.keep_with_next=True
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),hexc);tcPr.append(sh)
def P(text='',bold=False,color=None,size=None,italic=False,align=None,style=None):
    p=d.add_paragraph(style=style);r=p.add_run(text);r.bold=bold;r.italic=italic
    if color:r.font.color.rgb=color
    if size:r.font.size=Pt(size)
    if align:p.alignment=align
    return p
def H(t,l=1):return d.add_heading(t,l)
def bullets(items):
    for t in items:
        p=d.add_paragraph(style='List Bullet');p.add_run(t)
def table(rows,widths,fs=9.5,head=True):
    t=d.add_table(rows=len(rows),cols=len(rows[0]));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,r in enumerate(rows):
        for j,v in enumerate(r):
            c=t.cell(i,j);c.width=Cm(widths[j]);c.text='';p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
            run=p.add_run(str(v) if v is not None else '');run.font.size=Pt(fs)
            if i==0 and head:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);shade(c,'14213D')
            elif i%2==0:shade(c,'F3F6FA')
    for j,w in enumerate(widths): t.columns[j].width=Cm(w)
    if head:
        trPr=t.rows[0]._tr.get_or_add_trPr();h=OxmlElement('w:tblHeader');h.set(qn('w:val'),'true');trPr.append(h)
    d.add_paragraph().paragraph_format.space_after=Pt(2);return t
def box(label,text,kind):
    col={'quote':('EEF4FB',BLUE),'interp':('FEF3E6',ORG),'note':('F3F6FA',GRY)}[kind]
    t=d.add_table(rows=1,cols=1);t.style='Table Grid';t.autofit=False;t.columns[0].width=Cm(16.6);c=t.cell(0,0);shade(c,col[0]);c.width=Cm(16.6)
    p=c.paragraphs[0];r=p.add_run(label+'  ');r.bold=True;r.font.color.rgb=col[1];r.font.size=Pt(9.5)
    r2=p.add_run(text);r2.font.size=Pt(10)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
def quote(id_,note=''):
    q=Q[id_];box(f"원문 {id_} · {q['source']} {q['section']} · PDF p.{q['pages']}",'“'+' '.join(q['text'].split())+'”\n참고 번역(연구자): '+q['ko']+note,'quote')
# footer page numbers
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=fp.add_run('BCMS 연계성 해설서 V5 · 쪽 ');r.font.size=Pt(9)
for tag,txt in(('begin',None),(None,'PAGE'),('end',None)):
    rr=fp.add_run();rr.font.size=Pt(9)
    if tag:e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),tag);rr._r.append(e)
    else:e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=txt;rr._r.append(e)
# ---- title
P('BCMS 연계성 이야기하기',True,NAVY,26,align=WD_ALIGN_PARAGRAPH.LEFT).paragraph_format.space_before=Pt(90)
P('해설서 V5 — Arias의 22개 프로세스에서 연결품질 설문 문항까지',False,GRY,14)
P('박사논문 발표자료 짝 문서 · 연구 기준 2026년 10월 9일(V4를 시작점으로 보존)',False,GRY,11)
P('')
box('이 문서의 약속','모든 절의 끝에 출처를 적는다. 파란 상자는 원문 인용, 주황 상자는 연구자 해석·정의이다. 확인하지 못한 것은 확인하지 못했다고 쓴다. 짝 자료: BCMS_연계성_V5.xlsx(근거 시트), BCMS_연계성_이야기하기_V5.pptx(발표).','note')
d.add_page_break()
# ---- 읽는 법
H('0. 먼저 읽을 것: 인용을 읽는 법')
P('이 문서의 근거는 세 가지 확인 수준으로 구분한다.')
table([['수준','대상','의미와 한계'],
['① 문자열 대조','Arias·Białas 72개 인용','제공 PDF의 지정 쪽에서 공백·하이픈·따옴표 차이를 무시하고 대조하여 72/72 일치(2026-10-09, outputs/quote_verification.json). 해석·번역·측정의 타당성 증거는 아님.'],
['② 연구자 확인(V4 기록)','외부 학술 문헌 15건','본문 확인 6건(S01·S03·S04·S10·S13·S15), 초록·소개만 4건(S06·S07·S08·S14), 서지만 5건(S02·S05·S09, 표준 S11·S12). 이번 작업에서 재확인하지 못함.'],
['③ 연구자 해석·정의','관계 정규화, 개념 정의, 문항, 분류','근거가 아니라 검증 대상인 제안.']],[3.3,4.2,9.1])
P('출처: outputs/quote_verification.json; 시트 외부학술근거·참고문헌(V4 연구자 기록).',italic=True,color=GRY,size=9)
# ---- 1
H('1. 이야기의 줄거리')
P('BCMS(업무연속성경영시스템)의 프로세스는 서로 연결되어 작동한다. 이 연구는 Arias-Aranda 외(2026)가 정리한 22개 프로세스 유형을 객체로 삼아, 원문에서 서술된 연결을 31개 관계 기록으로 정리하고, 그 연결이 조직에서 어떤 상태로 유지되는지를 묻는 후보 문항을 만든다. 점수와 요인 구조는 아직 확정하지 않았으며 전문가 심사·인지면접·예비조사가 남아 있다.')
table([['장','이야기','핵심 메시지'] ]+[[c['no'],c['title'],c['message']] for c in ST['chapters'][1:7]],[1.0,3.4,12.2])
P('출처: Arias-Aranda 외(2026) Table 5·§4.3; 구성은 연구자.',italic=True,color=GRY,size=9)
# ---- 2
H('2. 왜 “연결”인가')
quote('A01');quote('A02')
box('연구자 해석','프로세스를 낱개로 평가하는 것과 별개로, 프로세스 사이를 오가는 산출물·기준이 후속 업무에 이어지는지를 묻는 질문이 필요하다. 두 원문은 “연결이 중요하다”는 근거이지 “연결품질의 차원”의 근거가 아니다.','interp')
H('이 연구가 묻는 것과 아직 말하지 않는 것',2)
table([['묻는다','아직 말하지 않는다'],['한 조직의 BCMS 단위, 최근 12개월 후보 기간의 운영 연결 상태 · 24개 운영 후보 연결 · 정합성·근거 추적·증빙·변경 전파(후보 개념) · 담당자 보고(핵심정보제공자)','확정된 4요인 주장 · 차원 평균·동일가중 지수 · 인과 주장 · 최종 문항 수와 N≥150 같은 고정 기준 · ORAS 척도 근거']],[8.3,8.3],10)
H('세 자료의 역할',2)
table([['자료','역할','한계'],['Arias-Aranda 외(2026)','22개 프로세스 유형(Table 5)과 서술된 연결(§4.3)','연결품질 척도를 검증한 연구가 아님'],['Białas(2010)','객체·관계·속성·역량질의의 표현 방식(§2 pp.2–3), 방법만 차용','게재지·권호·DOI 서지 보강 필요'],['외부 학술 문헌','추적성·데이터 품질·측정모형·척도 개발의 개념 근거','BCMS 적용 정의·문항은 연구자 확장']],[3.8,7.2,5.6])
P('출처: A01 Arias §2.1 PDF p.4; A02 §2.1 PDF p.5; Białas §2 pp.2–3; Noy & McGuinness(2001) Step 1 PDF pp.4–5; MacKenzie 외(2011) pp.302–303.',italic=True,color=GRY,size=9)
# ---- 3
H('3. 22개 객체')
P('객체의 범위는 Arias Table 5(PDF p.14)의 22개 BCMS 프로세스 유형이다. 22개 유형은 프로세스 유형이며 조직에서 실제 수행되는 프로세스 인스턴스와 구별한다. 전문가 지명률과 원자료 분류는 Arias pp.11–13의 참고 정보이다. CRM은 지명률 3%로 범주 3(기각)이지만 Table 5의 22개 유형에 포함되어 연구자가 범위에 둔다.')
table([['ID','한국어','Arias 원명','범주','PDCA','지명률','원자료 분류']]+[[o['ID'],o['한국어'],o['Arias Table 5 원명'],o['원문 범주'],o['PDCA'],str(o['전문가 지명률 %'])+'%',o['원자료 분류']] for o in B['객체22']],[1.3,3.0,5.2,1.3,1.4,1.4,3.0],8.5)
P('출처: Arias-Aranda 외(2026) Table 5, PDF p.14(이름·지명률 22/22 원문 대조 일치); 지명률·분류 pp.11–13.',italic=True,color=GRY,size=9)
# ---- 4
H('4. 연결 찾기: 원문에서 관계로')
H('4.1 L06 한 건으로 보는 과정',2)
quote('A14')
box('연구자 정규화','“guide”를 “BIA의 산출물이 전략·솔루션 선정의 입력·기준이 된다(informs)”로 해석 → 관계 L06: BIA → STR, 의미 “근거가 된다”, 등급 E1, 운영 후보. 주의: 승인된 업무연속성 요구사항은 BIA 산출물이며 REQ의 법규·계약 요구사항과 구별한다.','interp')
H('4.2 31개 기록의 분류와 등급',2)
cnt={};[cnt.__setitem__(l['분류'],cnt.get(l['분류'],0)+1) for l in LK]
g={};[g.__setitem__(l['V4 등급'],g.get(l['V4 등급'],0)+1) for l in LK]
table([['분류','수','처리'],['운영 후보',cnt['운영 후보'],'산출물·사용 목적·품질 개념 적합성 검토'],['위계 표시 L13·L14',cnt['위계 표시'],'위계만 표시, 운영 전달 품질 문항 제외'],['집단 진술 L27·L30',cnt['집단 진술'],'주석 보존, CORE는 23번째 객체가 아님'],['근거 보류 L11·L17·L29',cnt['근거 보류'],'추가 근거 전 문항 심사·투입 보류']],[5,1.6,10],10)
table([['등급','수','뜻(문헌 직접성)'],['E1',g['E1'],'끝점·의미를 원문이 직접 서술'],['E2',g['E2'],'끝점·방향·대상을 연구자가 정규화 (L10, L12, L16, L23–L27, L31)'],['E3',g['E3'],'직접 관계 근거 부족']],[2,1.6,13],10)
P('E1/E2는 문헌 직접성 분류이며 전문가 합의가 아니다. V3의 E1 중 L16, L23–L27, L31은 정규화가 필요하여 E2로 조정했다(수정이력 R10). 출처: 시트 관계근거31·온톨로지규칙; Arias §4.3 A18·A23·A29·A30.',italic=True,color=GRY,size=9)
H('4.3 22×22 관계 MAP',2)
d.add_picture(os.path.join(ROOT,'outputs/BCMS_관계MAP_22x22_V5.png'),width=Cm(16.6))
P('그림. 행 = 출발 유형, 열 = 도착 유형. 빈 셀은 미등록이며 관계 부재의 증거가 아니다. 출처: 객체 Arias Table 5 p.14; 관계 §4.3 pp.14–16(연구자 추출).',italic=True,color=GRY,size=9)
H('4.4 범위의 공백',2)
P('운영 후보 24개에 등장하는 유형은 22개 중 18개이다. POL·AWR·WARN·REC는 운영 연결 공백으로 표시하며 근거 보강 또는 범위 제한으로 다룬다. 이 18/22는 객체 등장 범위일 뿐 BCMS 전체 활동의 포괄률이 아니다.')
table([['ID','후보','내용','근거·사유','처리']]+[[r['ID'],r['관계 후보'],r['관계 내용'],r['근거·보류 사유'],r['다음 처리']] for r in B['보류및누락']],[1.2,2.6,3.6,6.6,2.6],8.5)
P('출처: 시트 보류및누락(V4); Arias §4.3 A23 p.15, A25 p.16; 수정이력 R16.',italic=True,color=GRY,size=9)
# ---- 5
H('5. 무엇을 묻는가: 개념 후보와 학술 근거')
P('여섯 개 개념은 후보이다. 연결품질 후보 4(CO·TR·EV·TM), 관리조건 후보 1(AC), 경로 후보 1(FC)로 역할을 구분했고 요인 수는 확정하지 않았다(수정이력 R02). 각 개념에 대해 문헌이 말하는 범위와 BCMS 적용의 한계를 함께 적는다.')
extra={'CO':'S03 연구자 본문 확인(Appendix D1) · S06 저자기관 초록만 확인','TR':'S01 연구자 본문 확인(§5.1 PDF p.4)','EV':'B17 제공 PDF 문자열 대조 일치(PDF p.5 §2.8)','AC':'A49·A50 문자열 대조 일치 · S07 출판사 초록만 확인','TM':'S04 연구자 본문 확인(저널 pp.214–215)','FC':'A23·A25·A29 문자열 대조 일치 · S08 출판사 소개만 확인, 직접 근거에서 제외'}
for k in ['CO','TR','EV','TM','AC','FC']:
    c=CQ[k];H(f"{k} {c['수정 개념']} — {c['검토 역할']}",2)
    table([['조작적 정의 후보(연구자)',c['조작적 정의 후보']],['출처·위치',c['출처·위치']],['확인 수준',extra[k]],['원문이 말하는 범위',c['원문이 말하는 범위']],['전이·한계',c['전이·한계']]],[3.6,13],9.5,head=False)
quote('B17')
box('연구자 해석','B17은 훈련 결과와 기록이 인증 과정의 증거로 쓰인다는 서술이다. 이를 모든 연결의 “증빙가능성” 차원으로 일반화한 것은 연구자 확장이며 TR과의 중복을 검토한다. (V4의 “B17 PDF p.6” 표기는 Białas PDF 5쪽 확인 결과 p.5로 정정함 — 수정이력 R26)','interp')
quote('A49')
P('출처: 시트 품질개념6·인용지도; Gotel & Finkelstein(1994) §5.1; Wang & Strong(1996) App. D1; Pipino 외(2002) pp.214–215; Zowghi & Gervasi(2003)·Bovens(2007) 초록; Argyris(1977).',italic=True,color=GRY,size=9)
# ---- 6
H('6. 문항과 응답 설계')
H('6.1 L06 한 연결에서 나오는 문항 후보',2)
table([['개념','L06 문항 후보 (V4)']]+[[CQ[k]['수정 개념'] if k in CQ else k,IT['L06-'+k]['V4 문항 후보']] for k in ['CO','TR','EV','TM','AC']],[3.4,13.2],9.5)
P('문항은 연구자 작성 후보이며 상태는 “미검토”이다. 연결의 내용과 품질 개념이 모두 적합한 곳에서만 문항을 만들고 전 관계×전 개념 배정은 하지 않는다.',size=10)
H('6.2 문항 은행 179의 구성',2)
c2={};[c2.__setitem__(r['역할·검토 상태'],c2.get(r['역할·검토 상태'],0)+1) for r in B['예비문항179']]
table([['검토 기록','수','다음 처리'],['연결품질 후보',c2['연결 품질 후보'],'4개 개념 적합성·중복 검토'],['관리조건 후보',c2['관리조건 후보'],'역할 구분 후 포함 여부 판단'],['개선 경로 후보',8,'직접 근거·경로·종료 기준 보강 전 개념 보류'],['구조·근거 보류',c2['구조·근거 보류'],'ID 보존, 현재 투입 제외'],['종합 7 + 응답 검토·준거 9',16,'보조 용도, 독립 외부타당도 근거로 자동 사용 금지']],[5,1.6,10],10)
P('합계 96+24+8+35+16=179. 179개는 검토 은행이며 최종 문항 수가 아니다(수정이력 R17). 출처: 시트 예비문항179; Boateng 외(2018) Table 1.',italic=True,color=GRY,size=9)
H('6.3 객체·관계·관측의 층 (Białas 방식 차용)',2)
table([['층','정의','L06 예']]+[[r['요소']+' · '+r['유형·질의'],r['정의·규칙'],r['예시·검증 조건']] for r in B['온톨로지규칙'][:7]],[4.2,7.4,5],9)
P('개념 온톨로지 단계이며 OWL 작성·추론·질의 테스트는 수행하지 않았다. CQ(역량질의)는 지식구조 질의이며 설문 문항과 같지 않다. 출처: Białas §2 pp.2–3(방법만 차용); Noy & McGuinness(2001) Step 1; 시트 온톨로지규칙.',italic=True,color=GRY,size=9)
H('6.4 응답 설계',2)
table([['항목','후보 규칙','운영·해석 조건']]+[[r['항목'],r['후보 규칙'],r['운영·해석 조건']] for r in B['응답설계']],[2.8,4,9.8],9)
P('출처: 시트 응답설계(V4); 수정이력 R13.',italic=True,color=GRY,size=9)
# ---- 7
H('7. 검증, 점수, 한계')
H('7.1 Gate 1–8',2)
table([['Gate','목적','다음 검토','완료 조건·주의']]+[[r['Gate'],r['목적'],r['다음 검토'],r['완료 조건·주의']] for r in B['검증계획']],[1.5,3.2,5.6,6.3],8.5)
H('7.2 점수를 아직 내지 않는 이유',2)
P('이전 버전의 차원 평균·동일가중 지수는 지표와 개념의 관계(반영적/형성적)가 정해지기 전에는 의미를 보증할 수 없다. 반영적 후보는 EFA/CFA, 형성적·복합 지표는 다른 모형 평가, 프로파일 보고 가운데 개념 정의에 맞게 선택한다. 개인 응답 수를 조직 표본 수로 대체하지 않는다. 출처: MacKenzie 외(2011) pp.302–303; 시트 응답설계·검증계획; 수정이력 R12·R14.')
H('7.3 열린 과제 (연구자 수행)',2)
bullets(['외부 문헌: S06·S07·S14 초록만 확인, S08 소개만 확인, S02·S05·S09·S11·S12 본문·표준 원문 확인 필요(이번 작업에서 재확인 못 함)','Białas 논문의 게재지·권호·DOI 등 서지 정보','ORAS 원저·판본·문항·사용 조건·한국어 적용 근거 (후보 문헌만 있으며 확정하지 못함)','연결품질 차원 수, 책임명확성의 역할, FC의 종료 기준·경로, 측정모형','POL·AWR·WARN·REC 공백과 추가 추출 후보 M01–M05, 24개 운영 후보에 대한 전문가 검토(Gate 2–3)'])
H('7.4 제3장 연구방법 서술안',2)
box('수행한 작업','본 연구는 Arias-Aranda 등(2026)의 22개 BCMS 프로세스를 내용 영역으로 설정하였다. 프로세스 설명에서 관계 근거를 추출하고, Białas의 온톨로지 표현방식을 참고하여 객체 유형, 산출물, 관계 의미와 출처를 구조화하였다. 운영 관계 후보와 포함 관계·집단 진술·근거 부족 관계를 구분하였다.','note')
box('예정된 검증','정합성, 근거 추적가능성, 연결 증빙가능성, 변경전파 적시성을 연결품질 후보 개념으로 검토하며, 책임명확성과 개선조치 완결성의 역할은 별도로 평가한다. 전문가 심사와 인지면접으로 문항을 수정·축소하고 적합한 측정모형을 검증한 뒤, ORAS와의 판별타당도 및 조직회복탄력성과의 구조적 관계를 검증한다.','note')
# ---- refs
d.add_page_break();H('참고문헌과 확인 수준')
for r in ST['refs']:
    p=d.add_paragraph();p.paragraph_format.left_indent=Cm(0.8);p.paragraph_format.first_line_indent=Cm(-0.8);p.paragraph_format.space_after=Pt(1)
    p.add_run(r['full']).font.size=Pt(9.5)
    p2=d.add_paragraph();p2.paragraph_format.left_indent=Cm(0.8);rr=p2.add_run('확인 수준: '+r['level']+('  ·  doi:'+r['doi'] if r['doi'] else '')+'  ·  용도: '+r['use']);rr.font.size=Pt(8.5);rr.font.color.rgb=GRY;rr.italic=True
# ---- appendix A
d.add_page_break();H('부록 A. 관계 기록 31개와 근거')
table([['ID','출발→도착','관계 의미','분류','V3→V4','인용(쪽)','원문 요지 · 연구자 정규화 · 주의']]+[[l['관계ID'],l['출발']+'→'+l['도착'],l['관계 의미'],l['분류'],l['V3 등급']+'→'+l['V4 등급'],(l['인용ID']+' p.'+Q[l['인용ID']]['pages']) if l['인용ID'] in Q else '—',f"{l['원문 요지']}\n[정규화] {l['연구자 정규화']}\n[주의] {l['수정·주의']}"] for l in LK],[1.0,1.9,2.0,1.6,1.5,1.7,6.9],7.5)
H('부록 B. 인용 원문 72개 (Arias A01–A54, Białas B01–B18)')
P('원문은 제공 PDF의 지정 쪽과 문자열 대조로 일치를 확인했다. 번역은 연구자의 참고 번역이다.',size=9.5)
table([['ID','쪽·절','제공 원문 / 참고 번역']]+[[r['인용ID'],f"{r['출처']} p.{r['PDF 쪽']} {r['절']}",f"{r['제공 원문']}\n→ {r['참고 번역']}"] for r in B['인용원문72']],[1.2,3.4,12],7.5)
H('부록 C. 수정 이력')
table([['ID','중요도','대상','V3 문제','수정','근거']]+[[r['ID'],r['중요도'],r['대상'],r['V3 문제'],r['V4 수정'],r['근거']] for r in B['수정이력']]+[['R21–R26','V5','발표용 정리','—','이야기하기·인용지도·참고문헌 추가(R21), 인용 확인 수준 3단계(R22), 72개 재대조(R23), 서지 표기(R24), 22×22 MAP 이미지(R25), B17 쪽 번호 p.6→p.5 정정(R26). 연구 분류·문항·정의는 V4와 동일.','시트 수정이력']],[1.3,1.4,2.2,3.8,5.6,2.3],7.5)
d.save(os.path.join(ROOT,'outputs/BCMS_연계성_해설서_V5.docx'));print('ok')
